# -*- coding: utf-8 -*-
"""聚焦排查：美加净 SKU 前导零在各数据源的存续情况。只读。"""
import os
from collections import Counter, defaultdict
import openpyxl

BASE = r'C:\Users\zhangyufei1\WorkBuddy\2026-06-30-09-24-40'

def col_index(hdr, *cands):
    for i, h in enumerate(hdr):
        hs = str(h or '').strip()
        for c in cands:
            if hs == c or (c in hs):
                return i
    return None

def dump(path, sheet, sku_cols, brand_cols=(), limit=8, label=''):
    print('\n' + '-' * 70)
    print(f'{label or sheet}  ({os.path.basename(path)} / {sheet})')
    print('-' * 70)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        print('  [no sheet]')
        return
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = next(it, None)
    print('  header:', list(hdr))
    i = col_index(hdr, *sku_cols)
    ib = col_index(hdr, *brand_cols) if brand_cols else None
    print(f'  sku_col_index={i} brand_col_index={ib}')
    if i is None:
        return
    types = Counter()
    lengths = Counter()
    by_brand = defaultdict(Counter)
    lead0 = Counter()
    samples_lead0 = []
    for r in it:
        if r is None or i >= len(r):
            continue
        v = r[i]
        if v is None or v == '':
            continue
        s = str(v)
        types[type(v).__name__] += 1
        lengths[len(s)] += 1
        if s.startswith('0'):
            lead0[len(s)] += 1
            if len(samples_lead0) < limit and len(samples_lead0) < 40:
                samples_lead0.append((s, type(v).__name__))
        if ib is not None and ib < len(r):
            by_brand[r[ib]][len(s)] += 1
    print('  类型:', dict(types))
    print('  长度:', dict(sorted(lengths.items())))
    print('  以0开头的条数:', sum(lead0.values()), '长度分布:', dict(lead0))
    print('  以0开头样例:', samples_lead0[:limit])
    for b, c in list(by_brand.items())[:15]:
        print(f'   品牌 {b}: {dict(sorted(c.items()))}')

P = lambda n: os.path.join(BASE, n)

print('=' * 70)
print('A. 订单明细（页面/评分主表）')
print('=' * 70)
dump(P('data.xlsx'), '订单明细', ['SKU编码'], ['品牌名称'])
# 美加净 单独看
wb = openpyxl.load_workbook(P('data.xlsx'), read_only=True, data_only=True)
ws = wb['订单明细']
it = ws.iter_rows(values_only=True)
hdr = next(it)
i_sku = col_index(hdr, 'SKU编码'); i_b = col_index(hdr, '品牌名称')
mj = set()
allsku = set()
for r in it:
    if r is None: continue
    allsku.add(str(r[i_sku]))
    if r[i_b] == '美加净':
        mj.add(str(r[i_sku]))
print('\n美加净 SKU 全量(%d):' % len(mj), sorted(mj)[:30])
print('美加净中以0开头:', sorted([x for x in mj if x.startswith('0')])[:30])
print('全表以0开头:', sorted([x for x in allsku if x.startswith('0')])[:30])

print('\n' + '=' * 70)
print('B. inventory.xlsx 分仓计划（下调预警 plan 数据源）')
print('=' * 70)
wb = openpyxl.load_workbook(P('inventory.xlsx'), read_only=True, data_only=True)
ws = wb['分仓计划']
rows = list(ws.iter_rows(values_only=True))
print('header:', rows[0])
print('第2行:', rows[1])
print('第3行:', rows[2])
# 两块：col0-3 (2026-08), col5-8 (2026-09)
for label, ci in [('2026-08 块 col0', 0), ('2026-09 块 col5', 5)]:
    types = Counter(); lengths = Counter(); lead = Counter(); samp = []
    for r in rows[2:]:
        if r is None or ci >= len(r): continue
        v = r[ci]
        if v is None or v == '': continue
        s = str(v)
        types[type(v).__name__] += 1
        lengths[len(s)] += 1
        if s.startswith('0'):
            lead[len(s)] += 1
            if len(samp) < 10: samp.append((s, type(v).__name__))
    print(f'\n  [{label}] 类型={dict(types)} 长度={dict(sorted(lengths.items()))} 以0开头={sum(lead.values())} {dict(lead)}')
    print('   以0开头样例:', samp)

print('\n' + '=' * 70)
print('C. inventory.xlsx 基础数据（产品主数据）')
print('=' * 70)
dump(P('inventory.xlsx'), '基础数据', ['产品编码'], ['品牌'])

print('\n' + '=' * 70)
print('D. 产品.xlsx 产品列表（权威主数据）')
print('=' * 70)
dump(P('产品.xlsx'), '产品列表', ['产品编码'], ['品牌'], limit=40)

print('\n' + '=' * 70)
print('E. 产品.xlsx 单价 / 淘汰品')
print('=' * 70)
dump(P('产品.xlsx'), '单价', ['物料编码'], [], limit=20)
dump(P('产品.xlsx'), '淘汰品', ['产品编码'], [], limit=20)
