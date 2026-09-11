# -*- coding: utf-8 -*-
"""排查 SKU 编码格式：美加净是否应为 5 位（带前导 0）。只读，不改任何数据。"""
import json, os, sys
from collections import Counter, defaultdict

BASE = r'C:\Users\zhangyufei1\WorkBuddy\2026-06-30-09-24-40'

def sec(t):
    print('\n' + '=' * 70)
    print(t)
    print('=' * 70)

# ---------- 1. inventory-plan.json（下调预警 d.sku 的来源） ----------
sec('1. inventory-plan.json  (plan data / down-warning source)')
p = os.path.join(BASE, 'inventory-plan.json')
print('exists', os.path.exists(p), 'size MB', round(os.path.getsize(p) / 1024 / 1024, 2) if os.path.exists(p) else '-')
if os.path.exists(p):
    with open(p, 'r', encoding='utf-8') as f:
        j = json.load(f)
    print('top keys', list(j.keys())[:20])

# ---------- 2. inventory.xlsx 里 SKU 编码的原生类型 ----------
sec('2. 原生 Excel 单元格类型（看前导零到底存不存在）')
try:
    import openpyxl
except ImportError:
    print('openpyxl missing')
    sys.exit(0)

def scan_xlsx(path, sheet_names=None, label=''):
    if not os.path.exists(path):
        print('  [skip] not exists', path)
        return
    print('\n--', label or os.path.basename(path))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print('  sheets:', wb.sheetnames)
    for sn in wb.sheetnames:
        if sheet_names and sn not in sheet_names:
            continue
        ws = wb[sn]
        rows = ws.iter_rows(values_only=True)
        hdr = next(rows, None)
        if not hdr:
            continue
        # 找 SKU 编码列 & 品牌列
        idx_sku = None
        idx_brand = None
        for i, h in enumerate(hdr):
            h = str(h or '')
            if idx_sku is None and ('SKU' in h.upper() and '编码' in h):
                idx_sku = i
            if idx_brand is None and ('品牌' in h):
                idx_brand = i
        print(f'  [{sn}] header={list(hdr)[:14]}')
        print(f'        idx_sku={idx_sku} idx_brand={idx_brand}')
        if idx_sku is None:
            continue
        type_cnt = Counter()
        brand_len = defaultdict(Counter)
        sample = []
        for r in rows:
            if idx_sku >= len(r):
                continue
            v = r[idx_sku]
            if v is None or v == '':
                continue
            type_cnt[type(v).__name__] += 1
            b = r[idx_brand] if (idx_brand is not None and idx_brand < len(r)) else '?'
            brand_len[b][len(str(v))] += 1
            if len(sample) < 6:
                sample.append((b, repr(v), type(v).__name__))
        print(f'        类型分布: {dict(type_cnt)}')
        print(f'        样例: {sample}')
        for b, c in list(brand_len.items())[:12]:
            print(f'        品牌 {b}: 长度分布 {dict(sorted(c.items()))}')

scan_xlsx(os.path.join(BASE, 'data.xlsx'), None, 'data.xlsx (order detail)')
scan_xlsx(os.path.join(BASE, 'inventory.xlsx'), None, 'inventory.xlsx')
scan_xlsx(os.path.join(BASE, '产品.xlsx'), None, '产品.xlsx (master)')

# ---------- 3. JSON 侧：data.json 的 SKU 字段 ----------
sec('3. JSON 侧 SKU 值形态')
for fn in ['data.json', 'inventory-core.json', 'inventory-master.json', 'inventory-plan.json']:
    fp = os.path.join(BASE, fn)
    if not os.path.exists(fp):
        print('  [skip]', fn)
        continue
    with open(fp, 'r', encoding='utf-8') as f:
        j = json.load(f)
    print(f'\n-- {fn}  top keys={list(j.keys())[:12]}')
    # 递归找所有名为 sku/materialCode 的值，统计长度
    cnt = Counter()
    samples = {}
    def walk(o, depth=0):
        if depth > 4:
            return
        if isinstance(o, dict):
            for k, v in o.items():
                if k in ('sku', 'materialCode', 'skuCode', 'sku_code', '物料编码') and isinstance(v, (str, int)):
                    s = str(v)
                    cnt[len(s)] += 1
                    if len(s) not in samples:
                        samples[len(s)] = []
                    if len(samples[len(s)]) < 3:
                        samples[len(s)].append((repr(v), type(v).__name__))
                else:
                    walk(v, depth + 1)
        elif isinstance(o, list):
            for v in o[:400]:
                walk(v, depth + 1)
    walk(j)
    if cnt:
        print('   长度分布:', dict(sorted(cnt.items())))
        for L in sorted(samples):
            print(f'   len={L} 样例 {samples[L]}')
    else:
        print('   未发现 sku/materialCode 字段')
